from pathlib import Path
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
import csv

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".csv",
    ".xlsx",
}

def is_supported(path: str) -> bool:
    extension = Path(path).suffix.lower()

    return extension in SUPPORTED_EXTENSIONS


def extract_text(path: str) -> str:
    file_path = Path(path).expanduser().resolve()
    extension = file_path.suffix.lower()
    
    try:
        if extension == ".pdf":
            return extract_pdf(file_path)

        elif extension == ".docx":
            return extract_docx(file_path)

        elif extension == ".txt":
            return extract_txt(file_path)

        elif extension == ".csv":
            return extract_csv(file_path)

        elif extension == ".xlsx":
            return extract_xlsx(file_path)

        else:
            return ""

    except Exception as error:
        print(f"Could not read {path}: {error}")
        return ""


def extract_pdf(path: Path) -> str:
    reader = PdfReader(path)

    text = []

    for page in reader.pages:
        page_text = page.extract_text()

        if page_text:
            text.append(page_text)

    return "\n".join(text)


def extract_docx(path: Path) -> str:
    document = Document(path)

    paragraphs = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            paragraphs.append(paragraph.text)

    return "\n".join(paragraphs)


def extract_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def extract_csv(path: Path) -> str:
    rows = []

    with open(path, "r", encoding="utf-8", errors="ignore") as file:
        reader = csv.reader(file)

        for row in reader:
            rows.append(" | ".join(row))

    return "\n".join(rows)


def extract_xlsx(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)

    rows = []

    for sheet in workbook.worksheets:
        rows.append(f"--- Sheet: {sheet.title} ---")

        for row in sheet.iter_rows(values_only=True):
            values = [
                str(value)
                for value in row
                if value is not None
            ]

            if values:
                rows.append(" | ".join(values))

    return "\n".join(rows)